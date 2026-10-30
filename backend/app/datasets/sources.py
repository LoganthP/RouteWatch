"""Format readers. Every reader yields plain dict rows; point geometry is exposed as __lng / __lat.
CSV, TSV, TXT and ZIP entries stream; Excel, JSON, GeoJSON, KML, GPX and PDF are parsed as a whole (size-capped)."""
import csv
import io
import json
import math
import re
import sys
import zipfile
from datetime import date, datetime

from defusedxml import ElementTree as ET

from .. import config as cfg
from ..util import HttpError
from .files import ARCHIVE_EXT, IN_MEMORY, kind_of

MB = 1024 * 1024
csv.field_size_limit(2 ** 31 - 1 if sys.platform == "win32" else sys.maxsize)
COORD_LINE = re.compile(r"(-?\d{1,2}\.\d{3,})\s*[,;\s]\s*(-?\d{1,3}\.\d{3,})")


def user_error(msg):
    return HttpError(400, msg)


class _ZipEntry(io.RawIOBase):
    """Binary stream of one zip entry; closing it also closes the archive."""
    def __init__(self, zip_path, name):
        super().__init__()
        self._z = zipfile.ZipFile(zip_path)
        self._f = self._z.open(name)

    def readable(self):
        return True

    def readinto(self, b):
        data = self._f.read(len(b))
        b[:len(data)] = data
        return len(data)

    def close(self):
        try:
            self._f.close()
            self._z.close()
        finally:
            super().close()


class Source:
    def __init__(self, name, kind, size, path=None, zip_path=None):
        self.name, self.kind, self.size, self.path, self.zip_path = name, kind, size, path, zip_path

    def open(self):
        if self.zip_path:
            return io.BufferedReader(_ZipEntry(self.zip_path, self.name), buffer_size=MB)
        return open(self.path, "rb")

    def read_all(self, cap_mb=None):
        cap = (cap_mb or cfg.LIMITS.max_inmemory_mb) * MB
        with self.open() as f:
            data = f.read(cap + 1)
        if len(data) > cap:
            raise user_error(f"File is larger than {cap // MB} MB for this format. Convert it to CSV.")
        return data


def list_zip(zip_path):
    L, entries, warnings, total = cfg.LIMITS, [], [], 0
    try:
        zf = zipfile.ZipFile(zip_path)
    except zipfile.BadZipFile as e:
        raise user_error(f"Could not open the zip: {e}")
    with zf:
        count = 0
        for info in zf.infolist():
            name = info.filename
            if name.endswith("/") or any(p.startswith(".") or p == "__MACOSX" for p in name.split("/")):
                continue
            if ".." in name or name.startswith("/") or re.match(r"^[a-zA-Z]:", name):
                raise user_error(f"Unsafe path inside zip: {name}")
            count += 1
            if count > L.max_zip_entries:
                raise user_error(f"The zip has more than {L.max_zip_entries} files.")
            total += info.file_size
            if total > L.max_uncompressed_mb * MB:
                raise user_error(f"The zip expands to more than {int(L.max_uncompressed_mb)} MB.")
            kind = kind_of(name)
            if ARCHIVE_EXT.search(name):
                warnings.append(f"Skipped nested archive {name}.")
            elif not kind or kind == "zip":
                warnings.append(f"Skipped unsupported file {name}.")
            elif kind == "xls":
                warnings.append(f"Skipped {name}: legacy .xls is not supported, save it as .xlsx or .csv.")
            elif kind in IN_MEMORY and info.file_size > L.max_inmemory_mb * MB:
                warnings.append(f"Skipped {name}: {kind} files over {int(L.max_inmemory_mb)} MB are not supported, convert to CSV.")
            else:
                entries.append(Source(name, kind, info.file_size, zip_path=zip_path))
    return entries, warnings


def describe(path, file_name, kind, size):
    """→ (sources, warnings) for an uploaded file."""
    if kind == "zip":
        sources, warnings = list_zip(path)
        if not sources:
            raise user_error(f"The zip has no supported files. {' '.join(warnings)}".strip())
        return sources, warnings
    if kind == "xls":
        raise user_error("Legacy .xls files are not supported. Open it in Excel and save as .xlsx or .csv.")
    if kind in IN_MEMORY and size > cfg.LIMITS.max_inmemory_mb * MB:
        raise user_error(f"{kind.upper()} files over {int(cfg.LIMITS.max_inmemory_mb)} MB are not supported. Convert it to CSV, which streams without a size limit.")
    return [Source(file_name, kind, size, path=path)], []


# ---------------------------------------------------------------- delimited text
def _encoding(head):
    if head[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return "utf-16"
    try:
        head.decode("utf-8")
        return "utf-8-sig"
    except UnicodeDecodeError as e:
        return "utf-8-sig" if e.start >= len(head) - 4 else "cp1252"  # cut in the middle of a character is fine


def _sniff(text):
    lines = [l for l in text.splitlines() if l.strip()][:20]
    best, best_score = None, 0
    for d in (",", "\t", ";", "|"):
        counts = [l.count(d) for l in lines]
        if not counts:
            continue
        lo = min(counts)
        score = lo * 10 - (max(counts) - lo) if lo > 0 else 0
        if score > best_score:
            best, best_score = d, score
    return best


def _unique_header(cells):
    seen, out = {}, []
    for i, h in enumerate(cells):
        name = (str(h).strip() or f"column_{i + 1}")
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        out.append(name)
    return out


def _delimited(src):
    with src.open() as f:
        head = f.read(65536)
    enc = _encoding(head)
    delim = "\t" if src.kind == "tsv" else _sniff(head.decode(enc if enc != "utf-8-sig" else "utf-8", "replace"))
    with io.TextIOWrapper(src.open(), encoding=enc, errors="replace", newline="") as text:
        if not delim:  # free text: one record per line that contains a "lat, lng" pair
            for line in text:
                m = COORD_LINE.search(line)
                if m:
                    yield {"line": line.strip()[:500], "__lat": m.group(1), "__lng": m.group(2)}
            return
        header = None
        for row in csv.reader(text, delimiter=delim, skipinitialspace=True):
            if not row or not any(c.strip() for c in row):
                continue
            if header is None:
                header = _unique_header(row)
                continue
            n = len(header)
            vals = [c.strip() for c in row[:n]]
            if len(vals) < n:
                vals += [""] * (n - len(vals))
            yield dict(zip(header, vals))


# ---------------------------------------------------------------- Excel
def _xlsx(src):
    from openpyxl import load_workbook
    stream = io.BytesIO(src.read_all()) if src.zip_path else src.path
    wb = load_workbook(stream, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            header, yielded = None, 0
            for vals in ws.iter_rows(values_only=True):
                vals = ["" if v is None else v for v in vals]
                if header is None:
                    if sum(1 for v in vals if str(v).strip()) < 2:
                        continue
                    header = _unique_header(vals)
                    continue
                if not any(str(v).strip() for v in vals):
                    continue
                yielded += 1
                yield {h: (vals[i] if i < len(vals) else "") for i, h in enumerate(header)}
            if yielded:
                return  # first sheet that has data
    finally:
        wb.close()


# ---------------------------------------------------------------- JSON / GeoJSON
def _centroid(geom):
    if not geom:
        return None
    if geom.get("type") == "Point":
        return geom.get("coordinates")
    pts = []
    def walk(c):
        if c and isinstance(c[0], (int, float)):
            pts.append(c)
        else:
            for x in c or []:
                walk(x)
    walk(geom.get("coordinates"))
    return [sum(p[0] for p in pts) / len(pts), sum(p[1] for p in pts) / len(pts)] if pts else None


def _flatten(obj, prefix="", out=None):
    out = {} if out is None else out
    for k, v in (obj or {}).items():
        key = f"{prefix}.{k}" if prefix else k
        if isinstance(v, dict) and len(v) < 30 and prefix.count(".") < 1:
            _flatten(v, key, out)
        else:
            out[key] = json.dumps(v) if isinstance(v, list) else v
    return out


def _json(src):
    raw = src.read_all().decode("utf-8-sig", "replace")
    try:
        data = json.loads(raw)
    except ValueError as e:
        try:  # JSON Lines
            data = [json.loads(l) for l in raw.splitlines() if l.strip()]
        except ValueError:
            raise user_error(f"Invalid JSON: {e}")
    feats = data.get("features") if isinstance(data, dict) and data.get("type") == "FeatureCollection" else None
    if feats is None and isinstance(data, dict) and data.get("type") == "Feature":
        feats = [data]
    if feats is not None:
        for f in feats:
            c = _centroid((f or {}).get("geometry"))
            yield {**_flatten((f or {}).get("properties") or {}), **({"__lng": c[0], "__lat": c[1]} if c else {})}
        return
    arr = data if isinstance(data, list) else None
    if arr is None and isinstance(data, dict):
        arr = next((v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)), None)
    if arr is None:
        raise user_error("The JSON does not contain a list of records or a GeoJSON FeatureCollection.")
    for r in arr:
        if isinstance(r, dict):
            yield _flatten(r)


# ---------------------------------------------------------------- KML / GPX
def _local(tag):
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _text(el):
    return (el.text or "").strip() if el is not None else ""


def _find(el, name):
    return next((c for c in el.iter() if _local(c.tag) == name), None)


def _parse_xml(src):
    try:
        return ET.fromstring(src.read_all())
    except Exception as e:
        raise user_error(f"Could not read the XML file: {str(e)[:120]}")


def _kml(src):
    root = _parse_xml(src)
    for pm in (e for e in root.iter() if _local(e.tag) == "Placemark"):
        row = {"__name": _text(_find(pm, "name")), "__desc": re.sub(r"<[^>]+>", " ", _text(_find(pm, "description"))).strip()}
        coords = _text(_find(pm, "coordinates")).split()
        if coords:
            parts = coords[0].split(",")
            try:
                row["__lng"], row["__lat"] = float(parts[0]), float(parts[1])
            except (ValueError, IndexError):
                pass
        when = _text(_find(pm, "when")) or _text(_find(pm, "begin"))
        if when:
            row["__time"] = when
        for el in pm.iter():
            if _local(el.tag) == "Data" and el.get("name"):
                row[el.get("name")] = _text(_find(el, "value"))
            elif _local(el.tag) == "SimpleData" and el.get("name"):
                row[el.get("name")] = _text(el)
        yield row


def _gpx(src):
    root = _parse_xml(src)
    for tag in ("wpt", "rtept", "trkpt"):
        for p in (e for e in root.iter() if _local(e.tag) == tag):
            try:
                lat, lng = float(p.get("lat")), float(p.get("lon"))
            except (TypeError, ValueError):
                continue
            row = {"__lat": lat, "__lng": lng, "__name": _text(_find(p, "name")), "__desc": _text(_find(p, "desc")) or _text(_find(p, "cmt")),
                   "__time": _text(_find(p, "time")), "__kind": tag}
            if _find(p, "type") is not None:
                row["type"] = _text(_find(p, "type"))
            yield row


# ---------------------------------------------------------------- PDF (text layer only)
_SPLITTERS = [re.compile(r"\t"), re.compile(r"\s*\|\s*"), re.compile(r"\s*,\s*"), re.compile(r"\s{2,}")]


def _pdf(src):
    from pypdf import PdfReader
    try:
        text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(src.read_all())).pages)
    except HttpError:
        raise
    except Exception as e:
        raise user_error(f"Could not read the PDF: {str(e)[:120]}")
    if len(re.sub(r"\s", "", text)) < 20:
        raise user_error("This PDF has no text layer (it is probably scanned images). Export the table as CSV or Excel instead.")
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    best = None
    for rx in _SPLITTERS:
        freq = {}
        for l in lines:
            n = len(rx.split(l))
            if n >= 3:
                freq[n] = freq.get(n, 0) + 1
        if freq:
            n, hits = max(freq.items(), key=lambda kv: kv[1])
            if hits > (best[2] if best else 0):
                best = (rx, n, hits)
    if best and best[2] >= 3:
        rx, n, _ = best
        rows = [c for c in (rx.split(l) for l in lines) if len(c) == n]
        h_idx = next((i for i, c in enumerate(rows) if sum(1 for x in c if re.search(r"[a-z]", x, re.I)) >= math.ceil(n / 2)), -1)
        header = _unique_header(rows[h_idx] if h_idx >= 0 else [f"column_{i + 1}" for i in range(n)])
        for c in rows[h_idx + 1 if h_idx >= 0 else 0:]:
            if [x.strip() for x in c] == [h.strip() for h in header]:
                continue  # header repeated on later pages
            yield dict(zip(header, c))
        return
    for l in lines:
        m = COORD_LINE.search(l)
        if m:
            yield {"line": l[:500], "__lat": m.group(1), "__lng": m.group(2)}


_READERS = {"csv": _delimited, "tsv": _delimited, "txt": _delimited, "xlsx": _xlsx, "json": _json, "geojson": _json, "kml": _kml, "gpx": _gpx, "pdf": _pdf}


def rows_of(src):
    reader = _READERS.get(src.kind)
    if not reader:
        raise user_error(f"Unsupported format: {src.kind}")
    return reader(src)


def _cell(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float) and not math.isfinite(v):
        return None
    if isinstance(v, int) and not isinstance(v, bool) and abs(v) > 2 ** 62:
        return str(v)
    return v if v is None or isinstance(v, (str, int, float, bool)) else str(v)


def preview(sources):
    """Columns, up to 20 sample rows and a row-count estimate."""
    sample, cols, used = [], {}, None
    for s in sources:
        it = rows_of(s)
        try:
            for row in it:
                sample.append({k: _cell(v) for k, v in row.items()})
                cols.update({k: None for k in row})
                if len(sample) >= 20:
                    break
        finally:
            it.close()
        if sample:
            used = s
            break
    if not sample:
        raise user_error("No readable rows were found. Check that the file has a header row and data.")
    estimated = None
    if used.kind in ("csv", "tsv", "txt"):
        with used.open() as f:
            head = f.read(262144)
        lines = head.count(b"\n")
        if lines > 1:
            per_row = len(head) / lines
            estimated = round(sum(s.size / per_row for s in sources if s.kind in ("csv", "tsv", "txt")))
    return {"columns": list(cols), "sample": sample, "estimatedRows": estimated, "previewSource": used.name}
